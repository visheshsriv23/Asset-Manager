import math
import os
import io
import openpyxl
import csv
from datetime import datetime
from openpyxl import Workbook
import shutil
import uuid
from fastapi import APIRouter, Depends, Query, HTTPException, status, UploadFile, File
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
from sqlalchemy import or_
from typing import Optional
from app.database import get_db
from app.models import Asset, User, AssignmentHistory,  Employee
from app.routers.auth import get_current_user
from app.schemas import PaginatedAssetsResponse, AssetCreateRequest, AssetResponse
from app.schemas import AssetDetailResponse, AssignmentHistoryItem, StatusUpdateRequest, ReassignRequest, AssetUpdateRequest
from datetime import datetime, timezone

router = APIRouter(prefix="/api/assets", tags=["Assets"])
UPLOAD_DIR = "uploads"
os.makedirs(UPLOAD_DIR, exist_ok=True)

@router.get("", response_model=PaginatedAssetsResponse)
def list_assets(
    page: int = Query(1, ge=1),
    page_size: int = Query(10, ge=1, le=100),
    q: Optional[str] = Query(None),
    status: Optional[str] = Query(None),
    asset_type: Optional[str] = Query(None, alias="type"),
    condition: Optional[str] = Query(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = db.query(Asset).options(joinedload(Asset.current_holder))

    if q:
        search_term = f"%{q.strip()}%"
        query = query.filter(
            or_(
                Asset.tag.ilike(search_term),
                Asset.make_model.ilike(search_term),
                Asset.serial_number.ilike(search_term),
            )
        )
    if status and status != "All":
        query = query.filter(Asset.status == status)

    if asset_type and asset_type != "All":
        query = query.filter(Asset.type == asset_type)

    if condition and condition != "All":
        query = query.filter(Asset.condition == condition)

    total = query.count()
    total_pages = math.ceil(total / page_size) if total > 0 else 1

    offset = (page - 1) * page_size
    items = query.order_by(Asset.id.desc()).offset(offset).limit(page_size).all()

    return PaginatedAssetsResponse(
        items=items,
        total=total,
        page=page,
        page_size=page_size,
        total_pages=total_pages,
    )

@router.post("", response_model=AssetResponse, status_code=status.HTTP_201_CREATED)
def create_asset(
    payload: AssetCreateRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if db.query(Asset).filter(Asset.tag == payload.tag.upper()).first():
        raise HTTPException(status_code=400, detail="Asset tag already exists.")
    if db.query(Asset).filter(Asset.serial_number == payload.serial_number.upper()).first():
        raise HTTPException(status_code=400, detail="Serial number already exists.")

    holder_id = payload.current_holder_id if payload.status == "Assigned" else None

    new_asset = Asset(
        tag=payload.tag.strip().upper(),
        type=payload.type,
        make_model=payload.make_model.strip(),
        serial_number=payload.serial_number.strip().upper(),
        configuration=payload.configuration,
        purchase_date=payload.purchase_date,
        vendor=payload.vendor.strip(),
        invoice_number=payload.invoice_number.strip(),
        cost=payload.cost,
        warranty_expiry=payload.warranty_expiry,
        location=payload.location or "Bengaluru HQ",
        condition=payload.condition,
        status=payload.status,
        current_holder_id=holder_id,
        invoice_file_url=payload.invoice_file_url,
    )
    db.add(new_asset)
    db.flush()

    action_text = "Asset registered"
    action_text = "Asset registered"
    if payload.status == "Assigned" and holder_id:
        emp = db.query(Employee).filter(Employee.id == holder_id).first()
        action_text = f"Assigned to {emp.name}" if emp else "Asset registered and assigned"
    elif payload.status == "Shipped to":
        action_text = f"Shipped to {new_asset.location}"

    history_entry = AssignmentHistory(
        asset_id=new_asset.id,
        employee_id=holder_id,
        action=action_text,
        date=datetime.now(timezone.utc),
        notes=payload.notes or "Initial registration into system",
    )
    db.add(history_entry)
    db.commit()
    db.refresh(new_asset)

    return new_asset

@router.get("/{asset_id}", response_model=AssetDetailResponse)
def get_asset_detail(
    asset_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    asset = (
        db.query(Asset)
        .options(joinedload(Asset.current_holder))
        .filter(Asset.id == asset_id)
        .first()
    )
    if not asset:
        raise HTTPException(status_code=404, detail="Asset not found")

    raw_history = (
        db.query(
            AssignmentHistory.id,
            AssignmentHistory.action,
            AssignmentHistory.date,
            AssignmentHistory.notes,
            Employee.name.label("employee_name"),
            Employee.employee_id.label("employee_code"),
        )
        .outerjoin(Employee, AssignmentHistory.employee_id == Employee.id)
        .filter(AssignmentHistory.asset_id == asset.id)
        .order_by(AssignmentHistory.date.desc())
        .all()
    )

    history_items = [
        AssignmentHistoryItem(
            id=h.id,
            action=h.action,
            date=h.date,
            notes=h.notes,
            employee_name=h.employee_name,
            employee_code=h.employee_code,
        )
        for h in raw_history
    ]

    res = AssetDetailResponse.from_orm(asset)
    res.history = history_items
    return res

@router.post("/{asset_id}/status")
def update_asset_status(
    asset_id: int,
    payload: StatusUpdateRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    asset = db.query(Asset).filter(Asset.id == asset_id).first()
    if not asset:
        raise HTTPException(status_code=404, detail="Asset not found")
    
    old_status = asset.status
    asset.status = payload.status

    if payload.status == "Shipped to":
        if not payload.destination or not payload.destination.strip():
            raise HTTPException(
                status_code=400,
                detail="Destination is required when status is 'Shipped to'."
            )
        dest = payload.destination.strip()
        asset.location = dest
        asset.current_holder_id = None
        action_label = f"Shipped to {dest}"
    elif payload.status in ["Ready to assign", "In repair", "Hardware issue", "Retired"]:
        asset.current_holder_id = None

    note_details = payload.notes.strip() if payload.notes else ""
    if payload.status == "Shipped to":
        note_details = f"Dispatched to: {payload.destination.strip()}. " + note_details

    log = AssignmentHistory(
        asset_id=asset.id,
        action=f"Status changed to {payload.status}",
        date=datetime.now(timezone.utc),
        notes=payload.notes or f"Changed from {old_status} to {payload.status}",
    )
    db.add(log)
    db.commit()
    db.refresh(asset)
    return {"message": "Status updated successfully"}

@router.post("/import/excel")
async def import_assets_excel(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if not (file.filename.endswith(".xlsx") or file.filename.endswith(".xls")):
        raise HTTPException(status_code=400, detail="Only .xlsx or .xls files are supported")

    content = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise HTTPException(status_code=400, detail="Excel file is empty or missing data rows")

    # Assuming headers match standard template
    imported_count = 0
    errors = []

    for idx, row in enumerate(rows[1:], start=2):
        if not row or not row[0]:  # Skip empty rows
            continue

        tag = str(row[0]).strip().upper()
        asset_type = str(row[1]).strip() if row[1] else "Laptop"
        make_model = str(row[2]).strip() if row[2] else "Standard Asset"
        serial = str(row[3]).strip().upper() if row[3] else f"SN-{tag}"

        # Avoid duplicates
        if db.query(Asset).filter(Asset.tag == tag).first():
            errors.append(f"Row {idx}: Tag {tag} already exists. Skipped.")
            continue

        new_asset = Asset(
            tag=tag,
            type=asset_type,
            make_model=make_model,
            serial_number=serial,
            status=str(row[4]).strip() if row[4] else "Ready to assign",
            condition=str(row[5]).strip() if row[5] else "Good",
            location=str(row[6]).strip() if row[6] else "Bengaluru HQ",
            vendor=str(row[8]).strip() if len(row) > 8 and row[8] else "Vendor",
            invoice_number=str(row[9]).strip() if len(row) > 9 and row[9] else "N/A",
            cost=float(row[10]) if len(row) > 10 and row[10] else 0.0,
        )
        db.add(new_asset)
        db.flush()

        # Audit history log
        history_entry = AssignmentHistory(
            asset_id=new_asset.id,
            action="Bulk imported from Excel",
            date=datetime.now(timezone.utc),
            notes=f"Uploaded via spreadsheet by {current_user.name}",
        )
        db.add(history_entry)
        imported_count += 1

    db.commit()
    return {"message": f"Successfully imported {imported_count} assets", "errors": errors}

@router.post("/{asset_id}/reassign")
def reassign_or_return_asset(
    asset_id: int,
    payload: ReassignRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    asset = db.query(Asset).filter(Asset.id == asset_id).first()
    if not asset:
        raise HTTPException(status_code=404, detail="Asset not found")

    action_date = payload.effective_date or datetime.now(timezone.utc)
    old_holder_name = asset.current_holder.name if asset.current_holder else "Stock"

    if payload.condition_on_return:
        asset.condition = payload.condition_on_return

    if payload.employee_id is None:
        asset.current_holder_id = None
        asset.status = "Ready to assign"

        log = AssignmentHistory(
            asset_id=asset.id,
            action="Returned to stock",
            date=action_date,
            notes=payload.notes or f"Returned from {old_holder_name}. Condition recorded as {asset.condition}.",
        )
        db.add(log)
        db.commit()
        return {"message": "Asset returned to stock successfully"}

    new_holder = db.query(Employee).filter(Employee.id == payload.employee_id).first()
    if not new_holder:
        raise HTTPException(status_code=404, detail="Target employee not found")

    asset.current_holder_id = new_holder.id
    asset.status = "Assigned"

    log = AssignmentHistory(
        asset_id=asset.id,
        employee_id=new_holder.id,
        action=f"Assigned to {new_holder.name}",
        date=action_date,
        notes=payload.notes or f"Transferred from {old_holder_name} to {new_holder.name} ({new_holder.employee_id})",
    )
    db.add(log)
    db.commit()
    return {"message": f"Asset assigned to {new_holder.name} successfully"}

@router.post("/upload-invoice")
async def upload_invoice(
    file: UploadFile = File(...),
    current_user: User = Depends(get_current_user),
):
    allowed_extensions = {".pdf", ".png", ".jpg", ".jpeg", ".doc", ".docx"}
    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in allowed_extensions:
        raise HTTPException(
            status_code=400,
            detail="Unsupported file format. Please upload PDF, PNG, JPG, or DOC.",
        )

    unique_filename = f"{uuid.uuid4().hex[:8]}_{file.filename}"
    file_path = os.path.join(UPLOAD_DIR, unique_filename)

    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
    return {
        "filename": file.filename,
        "url": f"/uploads/{unique_filename}",
        "size_kb": round(os.path.getsize(file_path) / 1024),
    }

@router.put("/{asset_id}", response_model=AssetResponse)
def update_asset(
    asset_id: int,
    payload: AssetUpdateRequest,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    asset = db.query(Asset).filter(Asset.id == asset_id).first()
    if not asset:
        raise HTTPException(status_code=404, detail="Asset not found")

    cleaned_serial = payload.serial_number.strip().upper()
    if cleaned_serial != asset.serial_number:
        duplicate = db.query(Asset).filter(Asset.serial_number == cleaned_serial, Asset.id != asset.id).first()
        if duplicate:
            raise HTTPException(status_code=400, detail="Serial number already assigned to another asset.")

    asset.type = payload.type
    asset.make_model = payload.make_model.strip()
    asset.serial_number = cleaned_serial
    asset.configuration = payload.configuration.strip()
    asset.purchase_date = payload.purchase_date
    asset.vendor = payload.vendor.strip()
    asset.invoice_number = payload.invoice_number.strip()
    asset.cost = payload.cost
    asset.condition = payload.condition
    asset.location = payload.location.strip() if payload.location else asset.location
    asset.warranty_expiry = payload.warranty_expiry
    if payload.invoice_file_url:
        asset.invoice_file_url = payload.invoice_file_url

    if payload.warranty_expiry:
        if isinstance(payload.warranty_expiry, str):
            asset.warranty_expiry = datetime.fromisoformat(payload.warranty_expiry.replace("Z", "+00:00"))
        else:
            asset.warranty_expiry = payload.warranty_expiry
    else:
        asset.warranty_expiry = asset.warranty_expiry or asset.purchase_date

    audit_note = payload.edit_reason.strip() if payload.edit_reason else "Asset details updated by administrator."
    log = AssignmentHistory(
        asset_id=asset.id,
        action="Asset details updated",
        date=datetime.now(timezone.utc),
        notes=audit_note,
    )
    db.add(log)
    db.commit()
    db.refresh(asset)
    return asset

@router.get("/export/excel")
def export_assets_excel(
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    assets = db.query(Asset).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Assets Inventory"

    # Header Row
    headers = [
        "Asset Tag", "Type", "Make & Model", "Serial Number",
        "Status", "Condition", "Location", "Holder",
        "Vendor", "Invoice No", "Cost", "Purchase Date", "Warranty Expiry"
    ]
    ws.append(headers)

    # Data Rows
    for a in assets:
        holder_name = a.current_holder.name if a.current_holder else "Unassigned"
        ws.append([
            a.tag,
            a.type,
            a.make_model,
            a.serial_number,
            a.status,
            a.condition,
            a.location,
            holder_name,
            a.vendor,
            a.invoice_number,
            a.cost,
            a.purchase_date.strftime("%Y-%m-%d") if a.purchase_date else "",
            a.warranty_expiry.strftime("%Y-%m-%d") if a.warranty_expiry else "",
        ])

    file_stream = io.BytesIO()
    wb.save(file_stream)
    file_stream.seek(0)

    filename = f"asset_inventory_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        file_stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )